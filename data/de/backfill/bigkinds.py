"""빅카인즈 엑셀 → 백필 작업 DB (S15P21E105-108).

주 단위로 받은 NewsResult_*.xlsx 를 스트리밍으로 읽어 sqlite 작업 DB 를 만든다.
분석제외 행과 URL 없는 행은 버리고, link 기준으로 중복을 제거한다.
이 DB 가 본문 채움 러너(fill_runner)의 입력이자 체크포인트다.
"""

import glob
import logging
import os
import re
import sqlite3

from openpyxl import load_workbook

from crawler.rss_collector import encode_link

logger = logging.getLogger(__name__)

XLSX_NAME = re.compile(r"NewsResult_\d{8}-\d{8}\.xlsx$")

# 전국일간지 12개사 → 내부 코드. article_id 해시 입력이라 바꾸면 안 된다
PUBLISHER_ID = {
    "경향신문": "khan", "국민일보": "kmib", "내일신문": "naeil", "동아일보": "donga",
    "문화일보": "munhwa", "서울신문": "seoul", "세계일보": "segye",
    "아시아투데이": "asiatoday", "조선일보": "chosun", "중앙일보": "joongang",
    "한겨레": "hani", "한국일보": "hankookilbo",
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS jobs (
    link TEXT PRIMARY KEY,          -- encode_link 적용본 (전 구간 표준 표현)
    company TEXT NOT NULL,
    publisher_id TEXT NOT NULL,
    title TEXT NOT NULL,
    day TEXT NOT NULL,              -- 엑셀 '일자' (YYYYMMDD)
    category_str TEXT,              -- 통합 분류1 원문 (예: 정치>행정)
    reporter TEXT,
    body200 TEXT,                   -- 빅카인즈 200자 본문 (실패 시 대체용)
    status TEXT NOT NULL DEFAULT 'todo',  -- todo / full / fallback
    article TEXT,
    published TEXT,                 -- 페이지에서 추출한 KST 'YYYY-MM-DD HH:MM:SS'
    fetched_at TEXT                 -- 처리 시각 (UTC ISO)
);
CREATE INDEX IF NOT EXISTS idx_jobs_status ON jobs (publisher_id, status);
"""


def load_jobs(xlsx_dir, db_path):
    """엑셀 폴더 전체를 작업 DB 로 적재한다. 재실행해도 기존 행은 건드리지 않는다."""
    conn = sqlite3.connect(db_path)
    conn.executescript(SCHEMA)
    stats = {"loaded": 0, "excluded": 0, "no_url": 0, "dup": 0, "unknown_pub": 0}
    files = sorted(p for p in glob.glob(os.path.join(xlsx_dir, "NewsResult_*.xlsx"))
                   if XLSX_NAME.search(os.path.basename(p)))
    for path in files:
        workbook = load_workbook(path, read_only=True)
        rows = workbook.active.iter_rows(min_row=2, values_only=True)
        for row in rows:
            # 컬럼: 0=뉴스식별자 1=일자 2=언론사 3=기고자 4=제목 5=통합분류1 ... 16=본문 17=URL 18=분석제외
            company = str(row[2] or "").strip()
            url = str(row[17] or "").strip()
            if str(row[18] or "").strip():
                stats["excluded"] += 1
                continue
            if not url.startswith("http"):
                stats["no_url"] += 1
                continue
            if company not in PUBLISHER_ID:
                stats["unknown_pub"] += 1
                continue
            inserted = conn.execute(
                "INSERT OR IGNORE INTO jobs (link, company, publisher_id, title, day,"
                " category_str, reporter, body200) VALUES (?,?,?,?,?,?,?,?)",
                (encode_link(url), company, PUBLISHER_ID[company],
                 str(row[4] or "").strip(), str(row[1] or "").strip(),
                 str(row[5] or "").strip() or None, str(row[3] or "").strip() or None,
                 str(row[16] or "").strip() or None),
            ).rowcount
            stats["loaded" if inserted else "dup"] += 1
        workbook.close()
        logger.info("%s 적재: %s", os.path.basename(path), stats)
    conn.commit()
    conn.close()
    return stats
